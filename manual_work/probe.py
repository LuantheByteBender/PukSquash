import sys, json, subprocess, shutil, winreg
from pathlib import Path
from docx import Document
root=Path.cwd(); ref=Path('C:/Users/300w12/Downloads/Campus Marketplace - Developer Manual (4).docx')
e=json.loads((root/'manual_work/reference_evidence.json').read_text())
print('SECTIONS',e['sections'])
doc=Document(ref)
for s in doc.styles:
    if s.name.lower() in ['normal','title','subtitle','heading 1','heading 2','heading 3']:
        print('STYLE',s.name,s.element.xml)
print('PARAS',[(p.style.name,p.text[:90],[(r.font.name,r.font.size.pt if r.font.size else None,r.bold) for r in p.runs[:2]]) for p in doc.paragraphs[:45]])
print('HEADINGS', '\n'.join(p.text for p in doc.paragraphs if p.style.name=='Heading 1'))
print('FOOTERS', e['footers'])
for name in ['Word.Application\\CLSID','LibreOffice.Application\\CLSID']:
    try:
        with winreg.OpenKey(winreg.HKEY_CLASSES_ROOT,name) as k: print(name,winreg.QueryValue(k,None))
    except OSError: pass
print('SOFFICE',shutil.which('soffice'))
for mod in ['win32com','comtypes','fitz']:
    try: __import__(mod); print('MODULE',mod,'OK')
    except ImportError: pass
print('GIT USER REPO HISTORY')
print(subprocess.run(['git','log','--oneline','-8','--','app/Repositories/UserRepository.php'],capture_output=True,text=True).stdout)
skill=Path('C:/Users/300w12/.codex/plugins/cache/openai-primary-runtime/documents/26.905.11957/skills/documents')
p=subprocess.run([sys.executable,str(skill/'render_docx.py'),str(ref),'--output_dir',str(root/'manual_work/reference_render'),'--emit_pdf'],capture_output=True,text=True)
(root/'manual_work/render_log.txt').write_text(p.stdout+p.stderr)
print('RENDER',p.returncode,(p.stdout+p.stderr)[-2200:])
