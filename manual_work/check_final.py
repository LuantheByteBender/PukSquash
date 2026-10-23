from pathlib import Path
import pypdfium2 as pdfium
from PIL import Image, ImageDraw
from zipfile import ZipFile,ZIP_DEFLATED
from docx import Document
import json, hashlib

root=Path.cwd();work=root/'manual_work';qa=work/'final_render';out=root/'output/PukSquash_Developer_Manual'
report=[]
for p in sorted(qa.glob('*.pdf')):
    pdf=pdfium.PdfDocument(str(p));folder=qa/p.stem;folder.mkdir(exist_ok=True)
    sheet=Image.new('RGB',(1220,1744),'#DDDDDD');stats=[]
    for n in range(len(pdf)):
        page=pdf[n];im=page.render(scale=1).to_pil().convert('RGB');im.save(folder/f'page-{n+1}.png')
        txt=page.get_textpage().get_text_range()
        stats.append({'page':n+1,'chars':len(txt),'first':txt[:100]})
        x=(n%2)*610;y=(n//2)*872
        if n<4:sheet.paste(im,(x,y+20));ImageDraw.Draw(sheet).text((x+10,y+3),f'{p.stem[:2]}  Page {n+1}',fill='black')
    sheet.save(qa/(p.stem+'_sheet.png'))
    report.append({'document':p.stem,'pages':len(pdf),'stats':stats})
    print(p.stem,len(pdf),'pages', [x['chars'] for x in stats])
(work/'final_qa.json').write_text(json.dumps(report,indent=2),encoding='utf-8')
assert len(report)==13
assert all(r['pages']==4 for r in report)
for p in out.glob('*.docx'):
    d=Document(p)
    for par in d.paragraphs:
        assert 'supabase' not in par.text.lower()
        assert 'campus-marketplace.git' not in par.text.lower()
    assert len(d.sections)==1
print('Structural checks passed; 52 pages ready for visual review.')
