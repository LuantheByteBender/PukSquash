from pathlib import Path
from docx import Document
from zipfile import ZipFile
from lxml import etree
import json, hashlib

base=Path.cwd()
out=base/'manual_work'
out.mkdir(exist_ok=True)
for parent in [*base.parents,base]:
    p=parent/'AGENTS.md'
    if p.exists(): print('AGENT INSTRUCTIONS',p,p.read_text(encoding='utf-8'))
files=[('business','Business_Case - PukSquash.docx'),('feasibility','Feasibility_Study - PukSquash.docx'),('functional','Functional Specifications PukSquash (3).docx'),('technical','Techinical specifications PukSquash (2).docx'),('developer_reference','Campus Marketplace - Developer Manual (4).docx'),('user_reference','Campus Marketplace - User Manual (3).docx')]
for key,name in files:
    path=Path('C:/Users/300w12/Downloads')/name
    doc=Document(path)
    lines=[]
    for el in doc.element.body:
        txt=' '.join(el.itertext()) if False else '\n'.join(''.join(p.itertext()) for p in [])
        if el.tag.endswith('}p'):
            from docx.text.paragraph import Paragraph
            p=Paragraph(el,doc)
            if p.text: lines.append(f'[{p.style.name}] {p.text}')
        elif el.tag.endswith('}tbl'):
            from docx.table import Table
            t=Table(el,doc)
            lines.extend(' | '.join(c.text for c in r.cells) for r in t.rows)
    (out/(key+'.txt')).write_text('\n'.join(lines),encoding='utf-8')
    print(key, 'paragraphs',len(doc.paragraphs),'tables',len(doc.tables),'sections',len(doc.sections),'words',sum(len(x.split()) for x in lines))
    if key=='developer_reference':
        ns={'w':'http://schemas.openxmlformats.org/wordprocessingml/2006/main'}
        evidence={'source':str(path),'sha256':hashlib.sha256(path.read_bytes()).hexdigest(),'sections':[{'width':s.page_width.inches,'height':s.page_height.inches,'top':s.top_margin.inches,'bottom':s.bottom_margin.inches,'left':s.left_margin.inches,'right':s.right_margin.inches} for s in doc.sections], 'styles':[{'name':s.name,'xml':s.element.xml} for s in doc.styles if s.name in ['Normal','Title','Subtitle','Heading 1','Heading 2','Heading 3']], 'paragraph_samples':[p._p.xml for p in doc.paragraphs[:35]], 'headers':[s.header._element.xml for s in doc.sections], 'footers':[s.footer._element.xml for s in doc.sections]}
        with ZipFile(path) as z:
            evidence['parts']={n:{'size':len(z.read(n)),'sha256':hashlib.sha256(z.read(n)).hexdigest()} for n in z.namelist()}
        (out/'reference_evidence.json').write_text(json.dumps(evidence,indent=2),encoding='utf-8')
print('EXTRACTION COMPLETE')
